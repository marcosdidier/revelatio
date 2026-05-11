"""Generate debtor_extraction.xlsx with a 36-column Excel Table named 'Debtors'.

The Power Automate Excel Online (Business) connector requires a *named Table* — not
just data in cells — so we create the file pre-formatted. Upload the resulting file
to SharePoint /Documents/Output/ and reference table 'Debtors' from the flow.

Run: .venv/bin/python 04_experiments/02_power_automate/setup/create_debtor_excel.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HEADERS = [
    # Admin (3) — populated by the flow itself, not by AI Builder
    "PDF_filename",
    "Data_processamento",
    "Confianca_media",
    # Grupo A — Identificação do Cliente (7, page 1)
    "Nome completo",
    "CPF",
    "Data de nascimento",
    "Telefone",
    "E-mail",
    "Endereço",
    "CEP",
    # Grupo B — Informações da Dívida (9, page 1)
    "Número do contrato",
    "Produto",
    "Data de contratação",
    "Valor original",
    "Saldo atualizado",
    "Dias em atraso",
    "Status interno",
    "Última tentativa de contato",
    "Responsável interno",
    # Grupo C — Comprovante de Endereço (8, page 2)
    "Titular do comprovante",
    "Endereço do comprovante",
    "Bairro/Cidade/UF",
    "CEP do comprovante",
    "Tipo de conta (Referência)",
    "Mês/Ano de referência",
    "Valor da conta",
    "Código de barras",
    # Grupo D — Comprovante Bancário (9, page 2)
    "Banco",
    "Agência",
    "Conta",
    "Data da transação",
    "Tipo de transação",
    "Pagador",
    "Recebedor",
    "Valor pago",
    "Código de autenticação",
]

assert len(HEADERS) == 36, f"expected 36 headers, got {len(HEADERS)}"

OUT = Path(__file__).resolve().parent / "debtor_extraction.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Debtors"

for col_idx, header in enumerate(HEADERS, start=1):
    ws.cell(row=1, column=col_idx, value=header)

# openpyxl Table requires at least one data row in the ref range, so seed row 2 blank.
# Power Automate "Add a row" appends at the end — first real row will land on row 3.
# After first real run completes, the user can manually delete row 2 if desired.
for col_idx in range(1, len(HEADERS) + 1):
    ws.cell(row=2, column=col_idx, value="")

last_col = get_column_letter(len(HEADERS))  # AJ for 36 columns
table_ref = f"A1:{last_col}2"

table = Table(displayName="Debtors", ref=table_ref)
table.tableStyleInfo = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)
ws.add_table(table)

# Reasonable column widths so the headers are readable when the analyst opens the file
for col_idx in range(1, len(HEADERS) + 1):
    ws.column_dimensions[get_column_letter(col_idx)].width = 22

wb.save(OUT)
print(f"✓ Wrote {OUT}")
print(f"  Sheet: '{ws.title}', Table: 'Debtors', Range: {table_ref}, Columns: {len(HEADERS)}")
